# -*- coding: utf-8 -*-
"""PPTX 문서의 공통 서식 - 팔레트 · 글꼴 · 상자 · 카드 · 제목줄.

매뉴얼과 분석 자료가 같은 모양이어야 해서 한 곳에 둔다. 예전에는
docs_make_manual.py 안에만 있어, 다른 문서를 만들면 팔레트를 베껴야 했다.
"""

import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- 팔레트
NAVY      = RGBColor(0x0F, 0x1B, 0x2D)   # 어두운 배경
NAVY_SOFT = RGBColor(0x1C, 0x2B, 0x45)   # 어두운 카드
EMBER     = RGBColor(0xE4, 0x57, 0x2E)   # 강조 (불꽃)
AMBER     = RGBColor(0xF4, 0xA2, 0x61)   # 보조 강조
TEAL      = RGBColor(0x2A, 0x9D, 0x8F)   # 정상/완료
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
PAPER     = RGBColor(0xF2, 0xF4, 0xF7)   # 밝은 카드
INK       = RGBColor(0x1A, 0x24, 0x35)   # 본문
MUTED     = RGBColor(0x66, 0x70, 0x85)   # 보조 텍스트
CODE_BG   = RGBColor(0x14, 0x20, 0x33)
CODE_FG   = RGBColor(0xD6, 0xDF, 0xEA)

KR   = "맑은 고딕"
MONO = "Consolas"

W, H = 13.333, 7.5
M    = 0.72          # 좌우 여백


def font_of(run, name=KR, size=14, bold=False, color=INK, italic=False):
    f = run.font
    f.name = name
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def textbox(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    tf._default_align = align          # 뒤에 추가되는 문단도 같은 정렬을 따르게 한다
    return tf


def para(tf, first=False, space_after=0, space_before=0, line=None, align=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align if align is not None else getattr(tf, "_default_align", PP_ALIGN.LEFT)
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if line:
        p.line_spacing = line
    return p


def write(tf, lines, first=True):
    """lines: [(text, dict(font kwargs), dict(para kwargs)), ...]"""
    is_first = first
    for item in lines:
        txt, fkw, pkw = item
        p = para(tf, first=is_first, **pkw)
        is_first = False
        r = p.add_run()
        r.text = txt
        font_of(r, **fkw)


def card(slide, x, y, w, h, fill=PAPER, radius=0.06, line=None, shadow=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def circle(slide, x, y, d, fill, text=None, tcolor=WHITE, tsize=16, tbold=True):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    if text is not None:
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        font_of(r, KR, tsize, tbold, tcolor)
    return sh


def dwidth(s):
    """고정폭 글꼴에서의 표시 너비. 한글·기호는 두 칸을 차지한다."""
    return sum(2 if ord(ch) > 0x2000 else 1 for ch in s)


def two_col(pairs, left_w=42):
    """
    모의 화면을 두 열로 그린다.

    공백 개수로 맞추면 한글이 섞인 줄에서 어긋난다. 한글을 두 칸으로
    세어 왼쪽 열을 같은 너비로 채운다.
    """
    out = []
    for l, r in pairs:
        if not r:
            out.append(l)
        else:
            out.append(l + " " * max(2, left_w - dwidth(l)) + r)
    return out


def code_block(slide, x, y, w, lines, size=12.5, pad=0.22):
    h = pad * 2 + len(lines) * (size * 1.55 / 72.0)
    card(slide, x, y, w, h, fill=CODE_BG, radius=0.05)
    tf = textbox(slide, x + pad, y + pad, w - pad * 2, h - pad * 2)
    items = []
    for i, ln in enumerate(lines):
        col = MUTED if ln.strip().startswith("#") else CODE_FG
        if ln.strip().startswith("$"):
            col = AMBER
        items.append((ln if ln else " ",
                      dict(name=MONO, size=size, color=col),
                      dict(space_after=1.5, line=1.0)))
    write(tf, items)
    return y + h


def slide_dark(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    return s


def slide_light(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    return s


def title_bar(slide, kicker, title, dark=False):
    """슬라이드 상단 제목 (장식선 없이 여백으로 구분)"""
    tf = textbox(slide, M, 0.52, W - 2 * M, 0.34)
    write(tf, [(kicker, dict(size=12.5, bold=True, color=EMBER),
                dict(space_after=0))])
    tf2 = textbox(slide, M, 0.88, W - 2 * M, 0.68)
    write(tf2, [(title, dict(size=30, bold=True, color=WHITE if dark else INK),
                 dict(space_after=0))])


def page_num(slide, n, dark=False):
    tf = textbox(slide, W - M - 1.0, H - 0.62, 1.0, 0.3, align=PP_ALIGN.RIGHT)
    write(tf, [(str(n), dict(size=10.5, color=MUTED), dict(space_after=0))])


